"""生成したPPTXを機械的に検査する（PowerPointルールの「ZIPとして全ファイルを検査」「元データとの機械比較」用）。

検査すること
1. [Content_Types].xml の Override がすべてZIP内に実在するか
2. 全スライドに発表者ノートがあるか（空でないか）
3. 和文フォントの文字セットが日本語（-128）になっているか、中国語（-122）が残っていないか
4. グラフに不正な線種（カンマ入りの prstDash）がないか
5. 図形がスライドの外にはみ出していないか（回転した線は両端の座標で判定）
6. 仮置きの文字（TODO・lorem など）が残っていないか
7. 外部へのハイパーリンクの数

使い方: python check_deck.py <対象.pptx> [期待するスライド枚数]
"""

import math
import re
import sys
import zipfile

from pptx import Presentation
from pptx.util import Emu

EMU_PER_INCH = 914400
TOLERANCE_INCH = 0.02
PLACEHOLDER_PATTERN = re.compile(r"TODO|lorem|ipsum|\[insert|xxx", re.I)


def check_content_types(zf):
    names = set(zf.namelist())
    xml = zf.read("[Content_Types].xml").decode("utf-8")
    missing = [p for p in re.findall(r'<Override PartName="/([^"]+)"', xml) if p not in names]
    return missing


def check_fonts_and_dash(zf):
    chinese = 0
    japanese = 0
    bad_dash = []
    for name in zf.namelist():
        if not re.match(r"ppt/(slides/slide\d+|charts/chart\d+)\.xml$", name):
            continue
        xml = zf.read(name).decode("utf-8")
        chinese += len(re.findall(r'charset="-122"', xml))
        japanese += len(re.findall(r'<a:ea typeface="BIZ UDGothic" charset="-128"/>', xml))
        if re.search(r'<a:prstDash val="[^"]*,[^"]*"', xml):
            bad_dash.append(name)
    return chinese, japanese, bad_dash


def shape_extents(shape):
    """図形の左右上下（インチ）。回転した線は両端の位置で求める"""
    x = shape.left / EMU_PER_INCH
    y = shape.top / EMU_PER_INCH
    w = shape.width / EMU_PER_INCH
    h = shape.height / EMU_PER_INCH
    rotation = getattr(shape, "rotation", 0) or 0
    if rotation and h < 1e-6:
        cx, cy = x + w / 2, y + h / 2
        a = math.radians(rotation)
        dx, dy = math.cos(a) * w / 2, math.sin(a) * w / 2
        xs = [cx - dx, cx + dx]
        ys = [cy - dy, cy + dy]
        return min(xs), max(xs), min(ys), max(ys)
    return x, x + w, y, y + h


def main(path, expected_slides=None):
    problems = []
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad:
            problems.append(f"ZIPの破損: {bad}")
        missing = check_content_types(zf)
        if missing:
            problems.append(f"Content_Typesに実在しない部品: {missing}")
        chinese, japanese, bad_dash = check_fonts_and_dash(zf)
        if chinese:
            problems.append(f"中国語の文字セットが {chinese} 箇所残っている")
        if bad_dash:
            problems.append(f"不正な線種: {bad_dash}")

    prs = Presentation(path)
    slide_w = prs.slide_width / EMU_PER_INCH
    slide_h = prs.slide_height / EMU_PER_INCH
    n_slides = len(prs.slides)
    if expected_slides is not None and n_slides != expected_slides:
        problems.append(f"スライド枚数が {n_slides} 枚（期待 {expected_slides} 枚）")

    notes_missing = []
    out_of_bounds = []
    placeholders = []
    links = 0
    shapes_total = 0
    for i, slide in enumerate(prs.slides, start=1):
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            notes_missing.append(i)
        for shape in slide.shapes:
            shapes_total += 1
            left, right, top, bottom = shape_extents(shape)
            if left < -TOLERANCE_INCH or top < -TOLERANCE_INCH or right > slide_w + TOLERANCE_INCH or bottom > slide_h + TOLERANCE_INCH:
                out_of_bounds.append((i, shape.shape_id, shape.name, round(left, 2), round(right, 2), round(top, 2), round(bottom, 2)))
            if shape.has_text_frame:
                text = shape.text_frame.text
                if PLACEHOLDER_PATTERN.search(text):
                    placeholders.append((i, text[:40]))
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.hyperlink and run.hyperlink.address:
                            links += 1
    if notes_missing:
        problems.append(f"ノートがないスライド: {notes_missing}")
    if out_of_bounds:
        problems.append(f"スライド外にはみ出す図形: {out_of_bounds}")
    if placeholders:
        problems.append(f"仮置きの文字: {placeholders}")

    print(f"スライド {n_slides} 枚 / 図形 {shapes_total} 個 / ノートあり {n_slides - len(notes_missing)} 枚 / 和文フォント指定 {japanese} 箇所 / 外部リンク {links} 個")
    if problems:
        print("問題あり:")
        for p in problems:
            print(" -", p)
        sys.exit(1)
    print("問題なし")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("使い方: python check_deck.py <対象.pptx> [期待するスライド枚数]")
    main(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else None)
